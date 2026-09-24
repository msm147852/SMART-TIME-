"""SMART AI-only external tools: PDF/Excel/Word/Search/Maps."""
import os, re
from pathlib import Path
from typing import List, Dict, Any
import pandas as pd
import requests
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from docx import Document

OUTPUT_DIR = Path(os.getenv("OPEN_MIND_OUTPUTS", Path(__file__).resolve().parents[2] / "outputs"))
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

def sanitize_filename(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|]', '', name)
    return (name.strip().replace(" ", "_")[:80] or "smart_ai_report")

def create_pdf(filename: str, title: str, content: str, table_data: List[List[str]] | None = None) -> str:
    path = OUTPUT_DIR / f"{sanitize_filename(filename)}.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    styles = getSampleStyleSheet(); story = [Paragraph(f"SMART AI - {title}", styles["Title"]), Spacer(1, 20)]
    for para in content.split("\n"):
        if para.strip(): story.extend([Paragraph(para, styles["Normal"]), Spacer(1, 10)])
    if table_data:
        t = Table(table_data); t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#111827")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), 1, colors.black)])); story.append(t)
    doc.build(story); return str(path)

def create_excel(filename: str, data: List[Dict[str, Any]]) -> str:
    path = OUTPUT_DIR / f"{sanitize_filename(filename)}.xlsx"; pd.DataFrame(data).to_excel(path, index=False); return str(path)

def create_word(filename: str, title: str, content: str) -> str:
    path = OUTPUT_DIR / f"{sanitize_filename(filename)}.docx"; doc = Document(); doc.add_heading(f"SMART AI - {title}", 0); doc.add_paragraph(content); doc.save(path); return str(path)

def web_search_brave(query: str, count: int = 8) -> Dict:
    key = os.getenv("BRAVE_SEARCH_API_KEY")
    if not key: return {"error": "BRAVE_SEARCH_API_KEY ناقص في .env - SMART AI", "results": []}
    r = requests.get("https://api.search.brave.com/res/v1/web/search", headers={"X-Subscription-Token": key}, params={"q": query, "count": count, "search_lang": "ar"}, timeout=15); r.raise_for_status(); data = r.json()
    return {"tool": "SMART AI Search", "query": query, "results": [{"title": i.get("title"), "url": i.get("url"), "desc": i.get("description")} for i in data.get("web", {}).get("results", [])]}

def google_maps_route(origin: str, destination: str) -> Dict:
    key = os.getenv("GOOGLE_MAPS_API_KEY")
    if not key: return {"error": "GOOGLE_MAPS_API_KEY ناقص - SMART AI", "no_invention": True}
    r = requests.get("https://maps.googleapis.com/maps/api/directions/json", params={"origin": origin, "destination": destination, "key": key, "language": "ar"}, timeout=15); r.raise_for_status(); data = r.json()
    if data.get("status") != "OK": return {"error": data.get("status"), "no_invention": True}
    leg = data["routes"][0]["legs"][0]
    return {"tool": "SMART AI Maps", "distance": leg["distance"]["text"], "duration": leg["duration"]["text"], "origin": leg["start_address"], "destination": leg["end_address"]}
